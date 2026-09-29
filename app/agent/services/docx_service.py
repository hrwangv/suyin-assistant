"""业务申请书 DOCX 渲染。

只有**一套模板**（templates/business_application.docx）：江苏省内 / 省外企业的文书结构
完全相同，差别只在其中一页的份数上。所以不再按地区分文件，改成渲染时算一个
`page_count` 传进模板（模板里可以用 {{ page_count }} 直接显示，或用循环重复那一页）。

渲染器优先用 docxtpl（模板里写 {{ company_name }}），
环境没装 docxtpl 时退回 python-docx 的占位符替换，保证功能可用。

【渲染器限制】python-docx 兜底只做 `{{ 字段 }}` 替换，**不认识 `{% for %}` 之类的
指令**。如果那一页要用循环重复 N 次，模板必须走 docxtpl；否则循环标签会被原样
留在文档里。届时应把 docxtpl 装成必装依赖（requirements.txt 里现在是可选项）。
"""
import copy
import re
import shutil
import tempfile
from datetime import date
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.conf.agent_config import agent_config
from app.core.logger import logger

JIANGSU = "jiangsu"
OTHER = "other"

# 唯一模板文件
TEMPLATE_FILE = "business_application.docx"
# 《综合信息查询授权书》的份数：省内一份、省外一式两份（备案要求）。
# 同一个模板靠这个值区分，改这里一处即可；也可由数据里的 page_count 覆盖。
PAGE_COUNT_BY_REGION = {JIANGSU: 1, OTHER: 2}
# 模板里需要按份数克隆的区块（用表格首格文字识别）
REPEATABLE_TABLE_KEY = "综合信息查询授权书"
# 副本是否另起一页：签署件一般一份一页，方便分别盖章
REPEAT_ON_NEW_PAGE = True


class TemplateNotFound(FileNotFoundError):
    """模板文件缺失。"""


def choose_template(province: str | None) -> tuple[str, str]:
    """按省份判定地区变体，返回 (template_type, template_path)。

    只有一套模板文件，所以 template_type 不再用来"选文件"，它现在的作用是
    决定那一页的份数（见 page_count_for），同时作为观测标签写进 State 与日志。
    """
    template_type = JIANGSU if _is_jiangsu(province) else OTHER
    template_path = agent_config.templates_dir / TEMPLATE_FILE
    if not template_path.exists():
        raise TemplateNotFound(f"未找到模板文件：{template_path}")
    return template_type, str(template_path)


def page_count_for(template_type: str, application_data: dict | None = None) -> int:
    """那一页要出几份：数据里显式给了就用给的，否则按地区变体取默认。"""
    override = (application_data or {}).get("page_count")
    if override not in (None, "", 0):
        try:
            return max(0, int(override))
        except (TypeError, ValueError):
            logger.warning(f"[docx] page_count 不是数字，忽略并改用默认值：{override!r}")
    return PAGE_COUNT_BY_REGION.get(template_type, PAGE_COUNT_BY_REGION[OTHER])


def _is_jiangsu(province: str | None) -> bool:
    text = (province or "").strip()
    return "江苏" in text


def _date_cn(iso_date: str) -> str:
    """2026-09-26 → 2026年09月26日（模板里的写法）。格式不对时原样返回。"""
    try:
        year, month, day = iso_date.split("-")
        return f"{int(year)}年{int(month):02d}月{int(day):02d}日"
    except (AttributeError, ValueError):
        return iso_date


def _template_with_repeats(template_path: str, page_count: int) -> tuple[str, int]:
    """按份数克隆可重复区块，返回 (可直接渲染的模板路径, 克隆了几份)。

    只在 page_count > 1 且模板里确实有该区块时才生成临时文件；
    其它情况原样返回，不产生额外 IO。
    """
    if page_count <= 1:
        return template_path, 0
    document = Document(template_path)
    clones = _clone_repeatable_tables(document, page_count)
    if not clones:
        logger.warning(
            f"[docx] 模板里没找到可重复区块「{REPEATABLE_TABLE_KEY}」，按一份处理"
        )
        return template_path, 0

    temp_dir = Path(tempfile.mkdtemp(prefix="docx_repeat_"))
    temp_path = temp_dir / Path(template_path).name
    document.save(str(temp_path))
    return str(temp_path), clones


def _clone_repeatable_tables(document, page_count: int) -> int:
    """把带识别关键词的表格克隆到 page_count 份（原表算一份）。"""
    clones = 0
    for table in list(document.tables):
        head = table.cell(0, 0).text or ""
        if REPEATABLE_TABLE_KEY not in head:
            continue
        anchor = table._tbl
        for _ in range(page_count - 1):
            clone = copy.deepcopy(anchor)
            if REPEAT_ON_NEW_PAGE:
                # 副本另起一页：先插分页符，再把副本接在后面
                break_paragraph = _page_break_paragraph()
                anchor.addnext(break_paragraph)
                break_paragraph.addnext(clone)
            else:
                anchor.addnext(clone)
            anchor = clone
            clones += 1
    return clones


def _page_break_paragraph():
    """造一个只含分页符的段落（python-docx 没有现成的 API）。"""
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    paragraph.append(run)
    return paragraph


def render_application(application_data: dict) -> dict:
    """渲染申请书，返回 {"file_path", "file_name", "url", "renderer", "page_count"}。

    步骤：选模板（只有一套）→ 算份数 → 需要时克隆可重复区块 → 渲染 → 落盘。
    """
    template_type, template_path = choose_template(application_data.get("province"))
    page_count = page_count_for(template_type, application_data)

    company_name = str(application_data.get("company_name") or "未知企业")
    application_date = str(application_data.get("application_date") or date.today().isoformat())
    file_name = f"{_safe(company_name)}_{application_date.replace('-', '')}_业务申请书.docx"

    agent_config.outputs_dir.mkdir(parents=True, exist_ok=True)
    output_path = agent_config.outputs_dir / file_name

    context = {
        **application_data,
        "template_type": template_type,
        "page_count": page_count,
        # 模板里的日期是中文写法（2026年02月25日）与年月（2026.02），单独给两个变量
        "application_date_cn": _date_cn(application_date),
        "application_month": application_date[:7].replace("-", "."),
    }

    working_template, clones = _template_with_repeats(template_path, page_count)
    try:
        renderer = _render_with_docxtpl(working_template, output_path, context)
        if renderer is None:
            renderer = _render_with_python_docx(working_template, output_path, context)
    finally:
        if working_template != template_path:
            shutil.rmtree(Path(working_template).parent, ignore_errors=True)

    if not output_path.exists():
        raise RuntimeError("DOCX 渲染失败：没有生成文件")

    url = f"{agent_config.file_url_prefix}/{quote(file_name)}"
    logger.info(
        f"[docx] 已生成申请书：{output_path}"
        f"（renderer={renderer}，variant={template_type}，"
        f"page_count={page_count}，重复区克隆={clones}）"
    )
    return {
        "file_path": str(output_path),
        "file_name": file_name,
        "url": url,
        "renderer": renderer,
        "template_type": template_type,
        "page_count": page_count,
    }


def _render_with_docxtpl(template_path: str, output_path: Path, context: dict):
    if agent_config.docx_renderer == "python-docx":
        return None
    try:
        from docxtpl import DocxTemplate
    except ImportError:
        return None
    try:
        doc = DocxTemplate(template_path)
        doc.render(context)
        doc.save(str(output_path))
        return "docxtpl"
    except Exception as exc:
        logger.warning(f"[docx] docxtpl 渲染失败，回退 python-docx：{exc}")
        return None


def _render_with_python_docx(template_path: str, output_path: Path, context: dict):
    """python-docx 兜底：把 {{ key }} 直接替换成文本。"""
    from docx import Document

    doc = Document(template_path)

    def replace_in_paragraph(paragraph):
        for run in paragraph.runs:
            if "{{" in run.text:
                run.text = _substitute(run.text, context)
        # 占位符被 Word 拆到多个 run 时整体重来一次
        if "{{" in paragraph.text:
            text = _substitute(paragraph.text, context)
            for run in paragraph.runs:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = text

    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)
    for section in doc.sections:
        for header_footer in (section.header, section.footer):
            for paragraph in header_footer.paragraphs:
                replace_in_paragraph(paragraph)

    doc.save(str(output_path))
    return "python-docx"


_PLACEHOLDER_RE = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


def _substitute(text: str, context: dict) -> str:
    def repl(match):
        value = context.get(match.group(1))
        return "" if value is None else str(value)

    return _PLACEHOLDER_RE.sub(repl, text)


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_") or "企业"
