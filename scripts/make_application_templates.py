"""生成业务申请书 DOCX 模板（只有一套）。

用法（项目根目录）：
    PYTHONPATH=. python scripts/make_application_templates.py

模板里用 {{ 字段名 }} 占位，docxtpl 和 python-docx 两种渲染器都能识别。
字段与 app/agent/schemas/application.py 的 ApplicationData 一致：
    company_name / unified_social_credit_code / province / city /
    registered_address / application_date / notes

江苏省内 / 省外的文书结构相同，差别只在其中一页的份数上：渲染时会额外传一个
page_count（见 docx_service），模板里可以 `{{ page_count }}` 直接显示，
或用 docxtpl 的循环重复那一页。
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.conf.agent_config import agent_config

TEMPLATE_FILE = "business_application.docx"


def build_template(path: Path) -> None:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("业 务 申 请 书")
    run.bold = True
    run.font.size = Pt(20)

    doc.add_paragraph()
    fields = [
        ("申请企业", "{{ company_name }}"),
        ("统一社会信用代码", "{{ unified_social_credit_code }}"),
        ("注册地址", "{{ registered_address }}"),
        ("所属省份", "{{ province }}"),
        ("所属城市", "{{ city }}"),
        ("申请日期", "{{ application_date }}"),
    ]
    for label, value in fields:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(value)

    doc.add_paragraph()
    doc.add_paragraph("尊敬的审批部门：").paragraph_format.first_line_indent = Pt(24)
    body = (
        "兹有我司向贵单位申请办理上述业务，特此提交本申请书。"
        "我司承诺所提交的全部材料真实、准确、完整，并愿意遵守相关业务规定。"
        "恳请贵单位予以审批。"
    )
    paragraph = doc.add_paragraph(body)
    paragraph.paragraph_format.first_line_indent = Pt(24)

    notes = doc.add_paragraph()
    notes.add_run("补充说明：").bold = True
    notes.add_run("{{ notes }}")
    notes.paragraph_format.first_line_indent = Pt(24)

    doc.add_paragraph()
    sign = doc.add_paragraph()
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sign.add_run("申请单位（盖章）：{{ company_name }}")
    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_line.add_run("日期：{{ application_date }}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def main() -> None:
    target = agent_config.templates_dir / TEMPLATE_FILE
    build_template(target)
    print(f"已生成模板：{target}")


if __name__ == "__main__":
    main()
